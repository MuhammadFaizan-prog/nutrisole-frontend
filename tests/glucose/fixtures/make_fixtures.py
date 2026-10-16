#!/usr/bin/env python3
"""Generates the SYNTHETIC glucose import fixtures (no real patient data, no real lab
branding). Layouts imitate common Pakistani lab reports and meter screens only in
structure. Re-run: python3 tests/glucose/fixtures/make_fixtures.py (needs Pillow,
openpyxl, reportlab). Outputs land next to this script; truth.json holds the expected
values used by tests/glucose/ocr-accuracy.mjs."""
import json, os, random, datetime
from PIL import Image, ImageDraw, ImageFont, ImageFilter
import openpyxl
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

HERE = os.path.dirname(os.path.abspath(__file__))
FONT = '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'
BOLD = '/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'
MONO = '/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf'
random.seed(7)

REPORTS = {
  'chughtai-style': {
    'lines': [
      ('CHUGHTAI-STYLE LAB REPORT  (SYNTHETIC SAMPLE - NOT A REAL REPORT)', True),
      ('Patient Name: Test Patient A          Age/Gender: 40 Y / M', False),
      ('Lab No: SYN-0001                      Collected On: 12-Oct-2026 08:20 AM', False),
      ('Registered On: 12-Oct-2026 08:05 AM   Reported On: 12-Oct-2026 01:10 PM', False),
      ('', False),
      ('Test Name                         Result     Unit      Reference Range', True),
      ('CHEMISTRY', True),
      ('Glucose Fasting                   112        mg/dL     70 - 100', False),
      ('HbA1c                             6.4        %         4.0 - 5.6', False),
      ('Estimated Average Glucose         137        mg/dL', False),
      ('', False),
      ('Interpretation: Fasting glucose 100-125 mg/dL impaired; >= 126 mg/dL diabetes.', False),
    ],
    'truth': {'readings': [{'value': '112', 'unit': 'mg/dL', 'context': 'fasting', 'date': 'Oct 12, 2026', 'time': '8:20 AM'}], 'nonGlucose': ['hba1c', 'eag']},
  },
  'shaukat-khanum-style': {
    'lines': [
      ('SKMCH-STYLE LABORATORY REPORT  (SYNTHETIC SAMPLE)', True),
      ('MR No: SYN-0002        Patient: Test Patient B', False),
      ('Specimen Collected: 14/10/2026 09:45      Reported: 14/10/2026 13:30', False),
      ('', False),
      ('Test                      Result    Units     Normal Range', True),
      ('Random Blood Sugar        168       mg/dl     < 200', False),
      ('Urine Sugar               Nil', False),
    ],
    'truth': {'readings': [{'value': '168', 'unit': 'mg/dL', 'context': 'random', 'date': 'Oct 14, 2026', 'time': '9:45 AM'}], 'nonGlucose': []},
  },
  'excel-labs-style': {
    'lines': [
      ('EXCEL-LABS-STYLE DIAGNOSTIC REPORT  (SYNTHETIC SAMPLE)', True),
      ('Patient: Test Patient C          Sample Collection Date: 15-10-2026 07:55', False),
      ('', False),
      ('Investigation            Result    Unit       Reference', True),
      ('FBS                      5.8       mmol/L     3.9 - 5.5', False),
      ('PPBS (2 Hrs)             9.1       mmol/L     < 7.8', False),
    ],
    'truth': {'readings': [{'value': '5.8', 'unit': 'mmol/L', 'context': 'fasting', 'date': 'Oct 15, 2026', 'time': '7:55 AM'}, {'value': '9.1', 'unit': 'mmol/L', 'context': 'after_meal', 'date': 'Oct 15, 2026', 'time': '7:55 AM'}], 'nonGlucose': []},
  },
  'idc-style': {
    'lines': [
      ('IDC-STYLE DIAGNOSTIC CENTRE REPORT  (SYNTHETIC SAMPLE)', True),
      ('Collection Date/Time : 16 Oct 2026 08:10 AM', False),
      ('Ref. By: Self            Patient: Test Patient D', False),
      ('', False),
      ('Parameter                         Result       Normal Values', True),
      ('Blood Glucose Random (RBS)        134 mg/dL    70 - 140', False),
      ('Glycated Hemoglobin (HbA1c)       7.1 %        < 5.7', False),
    ],
    'truth': {'readings': [{'value': '134', 'unit': 'mg/dL', 'context': 'random', 'date': 'Oct 16, 2026', 'time': '8:10 AM'}], 'nonGlucose': ['hba1c']},
  },
}
METERS = {
  'meter-mgdl': {'big': '126', 'unit': 'mg/dL', 'small': '16-10   07:42', 'truth': {'readings': [{'value': '126', 'unit': 'mg/dL', 'context': 'unspecified', 'date': 'Oct 16, 2026', 'time': '7:42 AM'}], 'nonGlucose': []}},
  'meter-mmol': {'big': '6.9', 'unit': 'mmol/L', 'small': '15/10   21:05', 'truth': {'readings': [{'value': '6.9', 'unit': 'mmol/L', 'context': 'unspecified', 'date': 'Oct 15, 2026', 'time': '9:05 PM'}], 'nonGlucose': []}},
}

def report_image(lines, path):
  img = Image.new('RGB', (1400, 160 + 52 * len(lines)), 'white'); d = ImageDraw.Draw(img)
  f, b = ImageFont.truetype(MONO, 24), ImageFont.truetype(BOLD, 24)
  y = 60
  for text, bold in lines:
    d.text((50, y), text, fill=(20, 20, 20), font=b if bold else f); y += 52
  d.line((50, 40, 1350, 40), fill=(120, 120, 120), width=2)
  img.save(path); return img

def degrade(img, path):
  img = img.rotate(1.5, expand=True, fillcolor='white', resample=Image.BICUBIC).filter(ImageFilter.GaussianBlur(1.1))
  px = img.load(); w, h = img.size
  for _ in range(w * h // 60):
    x, y = random.randrange(w), random.randrange(h); v = random.randint(140, 255); px[x, y] = (v, v, v)
  img = img.resize((int(w * 0.7), int(h * 0.7)))
  img.save(path, quality=45)

def meter_image(m, path):
  img = Image.new('RGB', (900, 640), (196, 208, 190)); d = ImageDraw.Draw(img)
  d.rounded_rectangle((40, 40, 860, 600), radius=40, outline=(60, 60, 60), width=6)
  d.text((90, 90), m['small'], fill=(25, 25, 25), font=ImageFont.truetype(MONO, 48))
  d.text((130, 190), m['big'], fill=(15, 15, 15), font=ImageFont.truetype(BOLD, 230))
  d.text((560, 470), m['unit'], fill=(25, 25, 25), font=ImageFont.truetype(BOLD, 56))
  img.save(path); return img

def report_pdf(lines, path, image=None):
  c = canvas.Canvas(path, pagesize=A4); c.setTitle('Synthetic glucose report fixture')
  if image:
    c.drawImage(image, 30, 400, width=535, height=400 * 0.9)
  else:
    y = 790
    for text, bold in lines:
      c.setFont('Courier-Bold' if bold else 'Courier', 9.5); c.drawString(40, y, text); y -= 18
  c.showPage(); c.save()

def main():
  truth = {}
  for name, r in REPORTS.items():
    img = report_image(r['lines'], f'{HERE}/{name}.png'); degrade(img, f'{HERE}/{name}-photo.jpg')
    truth[f'{name}.png'] = r['truth']; truth[f'{name}-photo.jpg'] = r['truth']
  for name, m in METERS.items():
    img = meter_image(m, f'{HERE}/{name}.png'); degrade(img, f'{HERE}/{name}-photo.jpg')
    truth[f'{name}.png'] = m['truth']; truth[f'{name}-photo.jpg'] = m['truth']
  report_pdf(REPORTS['chughtai-style']['lines'], f'{HERE}/chughtai-style.pdf')
  report_pdf(None, f'{HERE}/idc-style-scanned.pdf', image=f'{HERE}/idc-style.png')
  truth['chughtai-style.pdf'] = REPORTS['chughtai-style']['truth']; truth['idc-style-scanned.pdf'] = REPORTS['idc-style']['truth']
  # XLSX in a Contour-app-like shape with real date cells.
  wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Readings'
  ws.append(['Synthetic export - not real data']); ws.append([])
  ws.append(['Date and Time', 'Reading [mg/dL]', 'Meal Marker', 'Notes'])
  rows = [(datetime.datetime(2026, 10, 13, 7, 30), 104, 'Fasting', ''), (datetime.datetime(2026, 10, 13, 14, 5), 152, 'After Meal', 'lunch'), (datetime.datetime(2026, 10, 14, 7, 25), 98, 'Fasting', ''), (datetime.datetime(2026, 10, 14, 22, 40), 131, 'Night', 'bedtime check')]
  for r in rows:
    ws.append(list(r)); ws.cell(row=ws.max_row, column=1).number_format = 'yyyy-mm-dd hh:mm'
  wb.save(f'{HERE}/contour-style.xlsx')
  json.dump(truth, open(f'{HERE}/truth.json', 'w'), indent=1)

if __name__ == '__main__':
  main()
